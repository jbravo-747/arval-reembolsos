"""Genera los Excel del plan y los paquetes .zip de Power Automate (Fase 0).

Uso:
    python generar.py                      # IDs de Excel como marcadores (hay que reelegir archivo y tabla al importar)
    python generar.py <DRIVE_ID> <FILE_ID> # IDs reales del Inventario_correos.xlsx ya subido a OneDrive
"""
import json
import sys
import uuid
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = Path(__file__).parent / "salida"
ZONA = "Central Standard Time (Mexico)"

API = {
    "shared_office365": "Office 365 Outlook",
    "shared_excelonlinebusiness": "Excel Online (Business)",
}

INVENTARIO = ["Fecha", "Hora", "Remitente", "Dominio", "Asunto", "NumAdjuntos", "Adjuntos", "Extensiones"]
CLIENTES = ["Correo", "NombreCliente", "Notas", "Tipo", "Alias", "RFC"]
REGISTRO = [
    "ID", "FechaRecepcion", "HoraRecepcion", "IdCorreo", "Remitente", "Cliente", "Asunto",
    "NombreOriginal", "TamanoBytes", "NombreGuardado", "Ruta", "TipoDocumento", "UUID_CFDI",
    "RFC_Emisor", "Emisor", "Total", "FechaCFDI", "Duplicado", "IdOriginal", "Estado",
    "Confianza", "Observaciones", "EnResumen", "Enlace", "Prioridad",
]
MONITOR = ["Rutina", "Fecha", "Hora", "Resultado", "Detalle"]


# ---------- Excel ----------

def hoja_con_tabla(ws, encabezados, nombre_tabla):
    ws.append(encabezados)
    ultima = ws.cell(row=1, column=len(encabezados)).column_letter
    for col in range(1, len(encabezados) + 1):
        letra = ws.cell(row=1, column=col).column_letter
        ws.column_dimensions[letra].width = 22
        ws.column_dimensions[letra].number_format = "@"
        for fila in range(1, 3):
            ws.cell(row=fila, column=col).number_format = "@"
    tabla = Table(displayName=nombre_tabla, ref=f"A1:{ultima}2")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tabla)


def crear_excels():
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"
    hoja_con_tabla(ws, INVENTARIO, "tblInventario")
    wb.save(OUT / "Inventario_correos.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    hoja_con_tabla(ws, CLIENTES, "tblClientes")
    hoja_con_tabla(wb.create_sheet("Registro"), REGISTRO, "tblRegistro")
    hoja_con_tabla(wb.create_sheet("Monitor"), MONITOR, "tblMonitor")
    wb.save(OUT / "Control_Reembolsos.xlsx")


# ---------- Flujos ----------

def conector(api, operacion, parametros):
    return {
        "host": {
            "connectionName": api,
            "operationId": operacion,
            "apiId": f"/providers/Microsoft.PowerApps/apis/{api}",
        },
        "parameters": parametros,
        "authentication": "@parameters('$authentication')",
    }


def fila_inventario(drive, archivo, recibido, remitente, asunto, adjuntos):
    fecha = f"addHours({recibido}, -6)"
    return conector("shared_excelonlinebusiness", "AddRowV2", {
        "source": "me",
        "drive": drive,
        "file": archivo,
        "table": "tblInventario",
        "item/Fecha": f"@{{formatDateTime({fecha}, 'yyyy-MM-dd')}}",
        "item/Hora": f"@{{formatDateTime({fecha}, 'HH:mm')}}",
        "item/Remitente": f"@{{toLower({remitente})}}",
        "item/Dominio": f"@{{last(split(toLower({remitente}), '@'))}}",
        "item/Asunto": f"@{{{asunto}}}",
        "item/NumAdjuntos": f"@{{string(length(coalesce({adjuntos}, createArray())))}}",
        "item/Adjuntos": "@{join(body('Nombres'), ' | ')}",
        "item/Extensiones": "@{join(body('Exts'), ', ')}",
    })


def pasos_inventario(drive, archivo, recibido, remitente, asunto, adjuntos):
    desde = f"@coalesce({adjuntos}, createArray())"
    return {
        "Nombres": {
            "runAfter": {},
            "type": "Select",
            "inputs": {"from": desde, "select": "@item()?['name']"},
        },
        "Exts": {
            "runAfter": {"Nombres": ["Succeeded"]},
            "type": "Select",
            "inputs": {"from": desde, "select": "@toLower(last(split(item()?['name'], '.')))"},
        },
        "Agregar_fila": {
            "runAfter": {"Exts": ["Succeeded"]},
            "type": "OpenApiConnection",
            "inputs": fila_inventario(drive, archivo, recibido, remitente, asunto, adjuntos),
        },
    }


def definicion(triggers, actions):
    return {
        "$schema": "https://schema.management.azure.com/providers/Microsoft.Logic/schemas/2016-06-01/workflowdefinition.json#",
        "contentVersion": "1.0.0.0",
        "parameters": {
            "$connections": {"defaultValue": {}, "type": "Object"},
            "$authentication": {"defaultValue": {}, "type": "SecureObject"},
        },
        "triggers": triggers,
        "actions": actions,
    }


def flujo_inventario(drive, archivo):
    triggers = {
        "Cuando_llega_un_nuevo_correo": {
            "splitOn": "@triggerOutputs()?['body/value']",
            "type": "OpenApiConnectionNotification",
            "inputs": conector("shared_office365", "OnNewEmailV3", {
                "folderPath": "Inbox",
                "importance": "Any",
                "fetchOnlyWithAttachment": False,
                "includeAttachments": True,
            }),
            "runtimeConfiguration": {"concurrency": {"runs": 1}},
        }
    }
    actions = pasos_inventario(
        drive, archivo,
        "triggerOutputs()?['body/receivedDateTime']",
        "triggerOutputs()?['body/from']",
        "triggerOutputs()?['body/subject']",
        "triggerOutputs()?['body/attachments']",
    )
    return definicion(triggers, actions)


def flujo_retroactivo(drive, archivo):
    triggers = {
        "manual": {
            "type": "Request",
            "kind": "Button",
            "inputs": {"schema": {"type": "object", "properties": {}, "required": []}},
        }
    }
    actions = {
        "Correos": {
            "runAfter": {},
            "type": "OpenApiConnection",
            "inputs": conector("shared_office365", "GetEmailsV3", {
                "folderPath": "Inbox",
                "importance": "Any",
                "fetchOnlyUnread": False,
                "fetchOnlyWithAttachment": False,
                "includeAttachments": True,
                "top": 250,
            }),
        },
        "Desde_24sep": {
            "runAfter": {"Correos": ["Succeeded"]},
            "type": "Query",
            "inputs": {
                "from": "@body('Correos')?['value']",
                "where": "@greaterOrEquals(item()?['receivedDateTime'], '2026-09-24T06:00:00Z')",
            },
        },
        "Bucle": {
            "runAfter": {"Desde_24sep": ["Succeeded"]},
            "type": "Foreach",
            "foreach": "@body('Desde_24sep')",
            "actions": pasos_inventario(
                drive, archivo,
                "items('Bucle')?['receivedDateTime']",
                "items('Bucle')?['from']",
                "items('Bucle')?['subject']",
                "items('Bucle')?['attachments']",
            ),
            "runtimeConfiguration": {"concurrency": {"repetitions": 1}},
        },
    }
    return definicion(triggers, actions)


# ---------- Paquete (formato heredado de Importar paquete) ----------

def empaquetar(nombre, definicion_flujo, archivo_zip, conectores=API):
    id_flujo = str(uuid.uuid4())
    apis = {api: str(uuid.uuid4()) for api in conectores}
    conexiones = {api: str(uuid.uuid4()) for api in conectores}

    recursos = {
        id_flujo: {
            "type": "Microsoft.Flow/flows",
            "suggestedCreationType": "New",
            "creationType": "Existing, New, Update",
            "details": {"displayName": nombre},
            "configurableBy": "User",
            "hierarchy": "Root",
            "dependsOn": list(apis.values()) + list(conexiones.values()),
        }
    }
    for api, titulo in conectores.items():
        recursos[apis[api]] = {
            "id": f"/providers/Microsoft.PowerApps/apis/{api}",
            "name": api,
            "type": "Microsoft.PowerApps/apis",
            "suggestedCreationType": "Existing",
            "details": {"displayName": titulo},
            "configurableBy": "System",
            "hierarchy": "Child",
            "dependsOn": [],
        }
        recursos[conexiones[api]] = {
            "type": "Microsoft.PowerApps/apis/connections",
            "suggestedCreationType": "Existing",
            "creationType": "Existing",
            "details": {"displayName": titulo},
            "configurableBy": "User",
            "hierarchy": "Child",
            "dependsOn": [apis[api]],
        }

    manifiesto = {
        "schema": "1.0",
        "details": {
            "displayName": nombre,
            "description": "",
            "createdTime": "2026-09-29T00:00:00.0000000Z",
            "packageTelemetryId": str(uuid.uuid4()),
            "creator": "N/A",
            "sourceEnvironment": "",
        },
        "resources": recursos,
    }
    flujo = {
        "name": id_flujo,
        "id": f"/providers/Microsoft.Flow/flows/{id_flujo}",
        "type": "Microsoft.Flow/flows",
        "properties": {
            "apiId": "/providers/Microsoft.PowerApps/apis/shared_logicflows",
            "displayName": nombre,
            "definition": definicion_flujo,
            "connectionReferences": {
                api: {
                    "connectionName": "",
                    "source": "Embedded",
                    "id": f"/providers/Microsoft.PowerApps/apis/{api}",
                    "tier": "NotSpecified",
                }
                for api in conectores
            },
            "flowFailureAlertSubscribed": False,
            "isManaged": False,
        },
    }

    base = f"Microsoft.Flow/flows/{id_flujo}"
    with zipfile.ZipFile(archivo_zip, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("manifest.json", json.dumps(manifiesto, ensure_ascii=False))
        z.writestr("Microsoft.Flow/flows/manifest.json", json.dumps(
            {"packageSchemaVersion": "1.0", "flowAssets": {"assetPaths": [id_flujo]}}))
        z.writestr(f"{base}/definition.json", json.dumps(flujo, ensure_ascii=False))
        z.writestr(f"{base}/apisMap.json", json.dumps(apis))
        z.writestr(f"{base}/connectionsMap.json", json.dumps(conexiones))


def main():
    if len(sys.argv) > 1 and sys.argv[1] == "--solo-excel":
        OUT.mkdir(exist_ok=True)
        crear_excels()
        print("Excel generados en", OUT)
        return
    drive = sys.argv[1] if len(sys.argv) > 2 else "REELEGIR_BIBLIOTECA"
    archivo = sys.argv[2] if len(sys.argv) > 2 else "/Reembolsos/Inventario_correos.xlsx"
    OUT.mkdir(exist_ok=True)
    crear_excels()
    empaquetar("Reembolsos - 0 Inventario", flujo_inventario(drive, archivo),
               OUT / "Reembolsos_0_Inventario.zip")
    empaquetar("Reembolsos - 0b Inventario retroactivo", flujo_retroactivo(drive, archivo),
               OUT / "Reembolsos_0b_Inventario_retroactivo.zip")
    for f in sorted(OUT.iterdir()):
        print(f.name, f.stat().st_size)


if __name__ == "__main__":
    main()
