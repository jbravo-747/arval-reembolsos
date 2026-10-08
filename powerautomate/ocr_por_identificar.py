"""OCR local de la carpeta _Por identificar (rezago) con el OCR integrado de Windows.

Uso (desde Descargas\\Reembolsos_PowerAutomate):
    python ocr_por_identificar.py "<ruta local de la carpeta Reembolsos sincronizada>"            # solo reporte
    python ocr_por_identificar.py "<ruta local de la carpeta Reembolsos sincronizada>" --mover    # mueve lo identificado
    python ocr_por_identificar.py "<ruta ...>" --registro                                        # actualiza Cliente/Ruta en Registro

Qué hace:
- Lee el catálogo de la hoja Clientes de Control_Reembolsos.xlsx (NombreCliente, Alias, RFC; ignora Tipo Escaner/Aseguradora/Aviso).
- Para cada archivo de _Por identificar: XML → Receptor del CFDI; PDF → capa de texto si existe, si no OCR de las
  primeras páginas; imágenes → OCR. Busca nombres, alias y RFC del catálogo (sin acentos ni mayúsculas).
- Escribe OCR_resultado.csv junto a este script (archivo, cómo se leyó, cliente, confianza, acción). No guarda el texto leído.
- Con --mover: mueve a Clientes/<cliente>/<aaaa-mm> los archivos con un único cliente; los demás se quedan.
  Si un correo (mismo prefijo aaaammdd_hhmm) tiene un archivo identificado y otros sin identificar, mueve todos juntos.
- Con --registro: en la hoja Registro, para los archivos ya movidos, pone Cliente, Ruta, Estado=Revisar y una observación.
  Hazlo cuando el flujo no esté escribiendo (p. ej. fuera de horario) y con el archivo cerrado en Excel.
"""
import csv
import re
import shutil
import sys
import unicodedata
import xml.etree.ElementTree as ET
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import pymupdf as fitz  # PyMuPDF
from PIL import Image

AQUI = Path(__file__).parent
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # noqa: BLE001
    pass
PAGINAS_OCR = 3          # páginas por PDF que se leen con OCR
DPI = 200
EXT_IMG = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}
EXT_IGNORAR = {".htm", ".html", ".ini", ".ics", ".vcf"}
TIPOS_NO_CLIENTE = {"escaner", "aseguradora", "aviso", "ignorar"}


def normalizar(texto):
    t = unicodedata.normalize("NFKD", texto or "")
    t = "".join(c for c in t if not unicodedata.combining(c))
    t = re.sub(r"[^A-Za-z0-9 ]+", " ", t.upper())
    return re.sub(r"\s+", " ", t).strip()


def leer_catalogo(xlsx):
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["Clientes"]
    filas = ws.iter_rows(values_only=True)
    enc = [str(c or "").strip() for c in next(filas)]
    col = {n: i for i, n in enumerate(enc)}
    clientes = []
    for f in filas:
        if f is None or col.get("NombreCliente") is None:
            continue
        nombre = str(f[col["NombreCliente"]] or "").strip()
        if not nombre:
            continue
        tipo = str(f[col["Tipo"]] or "").strip().lower() if "Tipo" in col else ""
        if tipo in TIPOS_NO_CLIENTE:
            continue
        alias = str(f[col["Alias"]] or "").strip() if "Alias" in col else ""
        rfc = str(f[col["RFC"]] or "").strip().upper() if "RFC" in col else ""
        claves = {normalizar(nombre)}
        if alias:
            claves.add(normalizar(alias))
        clientes.append({"nombre": nombre, "claves": {c for c in claves if len(c) >= 6}, "rfc": rfc})
    wb.close()
    return clientes


def texto_pdf(ruta):
    """Devuelve (texto, origen). Usa la capa de texto si existe; si no, OCR de las primeras páginas."""
    doc = fitz.open(ruta)
    capa = " ".join(p.get_text() for p in doc)
    if len(normalizar(capa)) > 40:
        doc.close()
        return capa, "texto PDF"
    partes = []
    for p in list(doc)[:PAGINAS_OCR]:
        pix = p.get_pixmap(dpi=DPI)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        partes.append(ocr(img))
    doc.close()
    return " ".join(partes), "OCR"


def ocr(img):
    import winocr
    if max(img.size) > 9000:
        img.thumbnail((9000, 9000))
    return winocr.recognize_pil_sync(img, "es-MX").get("text", "")


def texto_xml(ruta):
    try:
        raiz = ET.parse(ruta).getroot()
    except ET.ParseError:
        return "", "XML ilegible"
    for el in raiz.iter():
        if el.tag.split("}")[-1] == "Receptor":
            return f"{el.get('Nombre', '')} {el.get('Rfc', '')}", "Receptor CFDI"
    return "", "XML sin Receptor"


def coincidencias(texto, clientes):
    t = normalizar(texto)
    hallados = []
    for c in clientes:
        exacta = any(k in t for k in c["claves"]) or (c["rfc"] and c["rfc"] in t)
        if exacta:
            hallados.append((c["nombre"], "exacta"))
            continue
        # parcial: todas las palabras del nombre (de 3+ letras) aparecen en el texto
        palabras = [p for p in normalizar(c["nombre"]).split() if len(p) >= 3]
        if palabras and all(re.search(rf"\b{re.escape(p)}\b", t) for p in palabras):
            hallados.append((c["nombre"], "parcial"))
    return hallados


def analizar(base):
    carpeta = base / "_Por identificar"
    clientes = leer_catalogo(base / "Control_Reembolsos.xlsx")
    print(f"Catálogo: {len(clientes)} clientes con nombre. Carpeta: {carpeta}")
    resultados = []
    for ruta in sorted(carpeta.iterdir()):
        if not ruta.is_file():
            continue
        ext = ruta.suffix.lower()
        fila = {"archivo": ruta.name, "correo": ruta.name[:13], "origen": "", "cliente": "", "confianza": "",
                "candidatos": "", "accion": "se queda"}
        try:
            if ext in EXT_IGNORAR:
                fila["origen"] = "ignorado"
            elif ext == ".xml":
                texto, fila["origen"] = texto_xml(ruta)
            elif ext == ".pdf":
                texto, fila["origen"] = texto_pdf(ruta)
            elif ext in EXT_IMG:
                texto, fila["origen"] = ocr(Image.open(ruta).convert("RGB")), "OCR"
            else:
                fila["origen"] = "formato no leído"
                texto = ""
            if fila["origen"] not in ("ignorado", "formato no leído"):
                h = coincidencias(texto, clientes)
                exactas = [n for n, c in h if c == "exacta"]
                if len(exactas) == 1:
                    fila.update(cliente=exactas[0], confianza="exacta", accion="mover")
                elif len(exactas) > 1:
                    fila.update(candidatos=" | ".join(exactas), confianza="varios")
                elif len(h) == 1:
                    fila.update(cliente=h[0][0], confianza="parcial", accion="revisar")
                elif h:
                    fila.update(candidatos=" | ".join(n for n, _ in h), confianza="varios parciales")
        except Exception as e:  # noqa: BLE001
            fila["origen"] = f"error: {type(e).__name__}: {str(e)[:60]}"
        resultados.append(fila)
        print(f"  {fila['archivo'][:50]:50} {fila['origen']:14} {fila['cliente'] or fila['candidatos'] or '-'} {fila['confianza']}")

    # Propagar dentro del mismo correo: si un archivo quedó identificado (exacta) y los demás no, van juntos.
    por_correo = defaultdict(list)
    for f in resultados:
        por_correo[f["correo"]].append(f)
    for grupo in por_correo.values():
        nombres = {f["cliente"] for f in grupo if f["accion"] == "mover"}
        if len(nombres) == 1:
            nombre = nombres.pop()
            for f in grupo:
                if f["accion"] != "mover" and f["origen"] not in ("ignorado",) and not f["candidatos"]:
                    f.update(cliente=nombre, confianza="mismo correo", accion="mover")
    salida = AQUI / "OCR_resultado.csv"
    with salida.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(resultados[0].keys()) if resultados else ["archivo"])
        w.writeheader()
        w.writerows(resultados)
    n_mover = sum(1 for f in resultados if f["accion"] == "mover")
    n_rev = sum(1 for f in resultados if f["accion"] == "revisar")
    print(f"\nArchivos: {len(resultados)} | para mover: {n_mover} | revisar (parcial): {n_rev} | reporte: {salida}")
    return resultados


def mover(base, resultados):
    carpeta = base / "_Por identificar"
    movidos = []
    for f in resultados:
        if f["accion"] != "mover":
            continue
        mes = f"{f['archivo'][:4]}-{f['archivo'][4:6]}"
        destino = base / "Clientes" / f["cliente"] / mes
        destino.mkdir(parents=True, exist_ok=True)
        origen = carpeta / f["archivo"]
        if (destino / f["archivo"]).exists():
            print(f"  ya existe, no se mueve: {f['archivo']}")
            continue
        shutil.move(str(origen), str(destino / f["archivo"]))
        movidos.append((f["archivo"], f["cliente"], f"/Reembolsos/Clientes/{f['cliente']}/{mes}/{f['archivo']}"))
        print(f"  movido: {f['archivo'][:50]} -> Clientes/{f['cliente']}/{mes}")
    with (AQUI / "OCR_movidos.csv").open("a", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        for m in movidos:
            w.writerow([datetime.now().isoformat(timespec="seconds"), *m])
    print(f"Movidos: {len(movidos)} (registro en OCR_movidos.csv)")


def actualizar_registro(base):
    """Pone Cliente, Ruta, Estado y Observaciones en Registro para los archivos de OCR_movidos.csv."""
    from openpyxl import load_workbook
    movidos = {}
    with (AQUI / "OCR_movidos.csv").open(encoding="utf-8-sig") as fh:
        for fila in csv.reader(fh):
            movidos[fila[1]] = (fila[2], fila[3])
    xlsx = base / "Control_Reembolsos.xlsx"
    wb = load_workbook(xlsx)
    ws = wb["Registro"]
    enc = [str(c.value or "").strip() for c in ws[1]]
    col = {n: i + 1 for i, n in enumerate(enc)}
    n = 0
    for r in range(2, ws.max_row + 1):
        guardado = str(ws.cell(r, col["NombreGuardado"]).value or "")
        if guardado in movidos:
            cliente, ruta = movidos[guardado]
            ws.cell(r, col["Cliente"]).value = cliente
            ws.cell(r, col["Ruta"]).value = ruta
            ws.cell(r, col["Estado"]).value = "Revisar"
            obs = str(ws.cell(r, col["Observaciones"]).value or "")
            ws.cell(r, col["Observaciones"]).value = (obs + " " if obs else "") + "Cliente identificado por OCR local; verificar"
            n += 1
    wb.save(xlsx)
    print(f"Filas actualizadas en Registro: {n}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    base = Path(sys.argv[1])
    if "--registro" in sys.argv:
        actualizar_registro(base)
    else:
        res = analizar(base)
        if "--mover" in sys.argv:
            mover(base, res)
